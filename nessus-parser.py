# -*- coding: utf-8 -*-

"""
Nessus Parser Programı
Bu program, .nessus çıktılarından interaktif olarak rapor ve istatistik dosyaları üretir.
UTF-8 karakter kodlamasını kullanır.
"""

import os
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Renkli çıktı için colorama
import colorama
from colorama import Fore, Style

# colorama init
colorama.init(autoreset=True)

def parse_nessus_to_excel(nessus_file_path, output_excel_path="nessus_output.xlsx", severities=None):
    """
    .nessus dosyasını parse edip iki farklı Excel oluşturur:
      1) output_excel_path -> Zafiyet detayı raporu
      2) istatistik_<output_excel_path> -> İstatistik dosyası
    
    'severities' verilirse, sadece bu risk seviyelerine sahip zafiyetler rapora dahil edilir.
    Konsolda renkli çıktı verilir.
    """

    # Terminal renkleri (Risk bazlı)
    terminal_color_map = {
        "Critical": Fore.RED + Style.BRIGHT,          # Kırmızı parlak
        "High":     Fore.LIGHTRED_EX + Style.BRIGHT, # Açık kırmızı
        "Medium":   Fore.YELLOW + Style.BRIGHT,
        "Low":      Fore.GREEN + Style.BRIGHT,
        "Info":     Fore.CYAN + Style.BRIGHT,
        "None":     Fore.WHITE + Style.DIM
    }

    # Excel renkleri (hücre dolgu)
    risk_colors = {
        "Critical": "FF8B0000",  # Koyu kırmızı
        "High":     "FFFF0000",  # Parlak kırmızı
        "Medium":   "FFFFA500",  # Turuncu
        "Low":      "FF00FF00",  # Yeşil
        "Info":     "FF87CEEB",  # Açık mavi
        "None":     "FFFFFFFF",  # Beyaz
    }

    # Hangi risk seviyeleri filtrelenecek?
    if severities and len(severities) > 0:
        severities = [sev.strip() for sev in severities if sev.strip()]
        print(Fore.BLUE + Style.BRIGHT + f"[INFO] Filtrelenecek risk seviyeleri: {', '.join(severities)}" + Style.RESET_ALL)
    else:
        severities = None  # Filtre yoksa hepsi dahil
        print(Fore.BLUE + Style.BRIGHT + "[INFO] Tüm risk seviyeleri rapora dahil edilecek." + Style.RESET_ALL)

    print(Fore.BLUE + Style.BRIGHT + f"[INFO] Kaynak .nessus dosyası: {nessus_file_path}" + Style.RESET_ALL)

    # Dosya var mı, XML parse edilebilir mi?
    if not os.path.isfile(nessus_file_path):
        print(Fore.RED + f"[-] Hata: '{nessus_file_path}' dosyası bulunamadı." + Style.RESET_ALL)
        return

    try:
        tree = ET.parse(nessus_file_path)
        root = tree.getroot()
        print(Fore.GREEN + "[INFO] .nessus dosyası başarıyla okundu ve parse edildi." + Style.RESET_ALL)
    except Exception as e:
        print(Fore.RED + f"[-] Hata: {nessus_file_path} dosyası okunamadı veya XML parse edilemedi.\n{e}" + Style.RESET_ALL)
        return

    # 1) Rapor Excel (Detaylar)
    wb = Workbook()
    ws = wb.active
    ws.title = "Nessus Sonuçları"
    
    # Zafiyetleri raporlayacağımız kolon başlıkları
    headers = [
        "Durum",         # 1
        "Notlar",        # 2
        "Host",          # 3
        "Port",          # 4
        "Risk",          # 5
        "Zafiyet Adı",   # 6 (eski Plugin Name)
        "Description",   # 7
        "Solution",      # 8
        "See Also",      # 9
        "Plugin Output", # 10
        "CVE",           # 11
        "CVSS"           # 12
    ]
    ws.append(headers)

    risk_col_index = headers.index("Risk") + 1  # 5. sütun (1-based)
    
    # 2) İstatistik tutmak için yapı
    unique_hosts = set()
    total_vulns = 0
    # Her risk seviyesi için bulgu sayısı
    counts_by_severity = {
        "Critical": 0,
        "High": 0,
        "Medium": 0,
        "Low": 0,
        "Info": 0,
        "None": 0
    }
    # Farklı zafiyet adları (uniq)
    unique_vulns = set()
    # Risk seviyesine göre uniq zafiyet adları
    unique_vulns_by_severity = {
        "Critical": set(),
        "High": set(),
        "Medium": set(),
        "Low": set(),
        "Info": set(),
        "None": set()
    }

    row_counter = 2

    report = root.find("Report")
    if report is None:
        print(Fore.RED + "[-] Hata: .nessus dosyasında 'Report' düğümü bulunamadı!" + Style.RESET_ALL)
        return

    print(Fore.GREEN + Style.BRIGHT + "[INFO] Zafiyet bilgileri toplanıyor..." + Style.RESET_ALL)

    for report_host in report.findall("ReportHost"):
        # Host IP
        host_ip = report_host.get("name", "Unknown Host")

        # HostProperties ile kesin IP al
        host_properties = report_host.find("HostProperties")
        if host_properties is not None:
            for tag in host_properties.findall("tag"):
                if tag.attrib.get('name') == "host-ip":
                    host_ip = tag.text
                    break

        unique_hosts.add(host_ip)  # Bu host’u kaydedelim
        host_vuln_count = 0

        for report_item in report_host.findall("ReportItem"):
            port = report_item.attrib.get("port", "N/A")
            # Risk
            risk_node = report_item.find("risk_factor")
            risk = risk_node.text if risk_node is not None else "None"

            # Filtre (risk seviyesi)
            if severities and risk not in severities:
                continue

            # Zafiyet Adı (plugin_name)
            zafiyet_adi_node = report_item.find("plugin_name")
            zafiyet_adi = zafiyet_adi_node.text if zafiyet_adi_node is not None else "N/A"

            # Description
            description_node = report_item.find("description")
            description = description_node.text if description_node is not None else ""

            # Solution
            solution_node = report_item.find("solution")
            solution = solution_node.text if solution_node is not None else ""

            # See Also
            see_also_list = report_item.findall("see_also")
            see_also_joined = ", ".join(sa.text for sa in see_also_list if sa.text)

            # Plugin Output
            plugin_output_node = report_item.find("plugin_output")
            plugin_output = plugin_output_node.text if plugin_output_node is not None else ""

            # CVE
            cve_list = report_item.findall("cve")
            cve_joined = ", ".join(c.text for c in cve_list if c.text)

            # CVSS
            cvss2_node = report_item.find("cvss_base_score")
            cvss3_node = report_item.find("cvss3_base_score")
            cvss_score = ""
            if cvss3_node is not None and cvss3_node.text:
                cvss_score = cvss3_node.text
            elif cvss2_node is not None and cvss2_node.text:
                cvss_score = cvss2_node.text

            # Rapor Excel Satırı
            durum = ""
            notlar = ""

            row_data = [
                durum,
                notlar,
                host_ip,
                port,
                risk,
                zafiyet_adi,
                description,
                solution,
                see_also_joined,
                plugin_output,
                cve_joined,
                cvss_score
            ]
            ws.append(row_data)

            # Excel renklendirme
            cell_risk = ws.cell(row=row_counter, column=risk_col_index)
            fill_color = risk_colors.get(risk, "FFFFFFFF")
            cell_risk.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            row_counter += 1
            host_vuln_count += 1
            total_vulns += 1

            # İstatistik toplama
            counts_by_severity[risk] = counts_by_severity.get(risk, 0) + 1
            unique_vulns.add(zafiyet_adi)
            unique_vulns_by_severity[risk].add(zafiyet_adi)

            # Terminal çıktısı
            color_prefix = terminal_color_map.get(risk, Fore.WHITE)
            msg = f"[DETAIL] Host: {host_ip}, Port: {port}, Risk: {risk}, Zafiyet: {zafiyet_adi}"
            print(color_prefix + msg + Style.RESET_ALL)

        if host_vuln_count > 0:
            print(Fore.GREEN + f"[INFO] {host_ip} adresinde {host_vuln_count} adet zafiyet bulundu." + Style.RESET_ALL)

    # "Durum" sütunu (A sütunu) drop-down list
    dv_status = DataValidation(
        type="list",
        formula1='"Giderildi,Giderilemedi,Devam Ediyor,False Positive"',
        allow_blank=True
    )
    durum_column_letter = get_column_letter(1)  # A
    max_rows = 10000
    dv_range = f"{durum_column_letter}2:{durum_column_letter}{max_rows}"
    dv_status.add(dv_range)
    ws.add_data_validation(dv_status)

    # Kolon genişlikleri (opsiyonel)
    for col in ws.columns:
        max_length = 0
        column_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
        ws.column_dimensions[column_letter].width = min(max_length + 2, 80)

    # Ana Excel kaydet
    try:
        wb.save(output_excel_path)
        print(Fore.GREEN + f"[INFO] Ana rapor Excel oluşturuldu: {output_excel_path}" + Style.RESET_ALL)
    except Exception as e:
        print(Fore.RED + f"[-] Hata: Ana Excel kaydedilemedi.\n{e}" + Style.RESET_ALL)
        return

    # 3) İstatistik Excel oluşturma
    #    Çıktı dosyası: "istatistik_<output_excel_path>"
    istatistik_output_excel_path = f"istatistik_{output_excel_path}"
    wb_stats = Workbook()
    ws_stats = wb_stats.active
    ws_stats.title = "Istatistik"

    # Bazı satırlar halinde verileri yazalım
    row_num = 1

    # Toplam taranan sistem (host) sayısı
    ws_stats.cell(row=row_num, column=1, value="Toplam taranan sistem (host) sayısı:")
    ws_stats.cell(row=row_num, column=2, value=len(unique_hosts))
    row_num += 1

    # Toplam zafiyet sayısı
    ws_stats.cell(row=row_num, column=1, value="Toplam zafiyet (bulgu) sayısı:")
    ws_stats.cell(row=row_num, column=2, value=total_vulns)
    row_num += 1

    # Farklı (uniq) zafiyet adı sayısı
    ws_stats.cell(row=row_num, column=1, value="Toplam farklı zafiyet adı sayısı:")
    ws_stats.cell(row=row_num, column=2, value=len(unique_vulns))
    row_num += 2  # Bir satır boşluk

    # Her risk seviyesi için total bulgu sayısı
    ws_stats.cell(row=row_num, column=1, value="Risk Seviyesi")
    ws_stats.cell(row=row_num, column=2, value="Toplam Bulgu Adedi")
    ws_stats.cell(row=row_num, column=3, value="Farklı Zafiyet Adedi")
    row_num += 1

    for sev in ["Critical", "High", "Medium", "Low", "Info", "None"]:
        bulgu_sayisi = counts_by_severity[sev]
        uniq_zafiyet_sayisi = len(unique_vulns_by_severity[sev])
        ws_stats.cell(row=row_num, column=1, value=sev)
        ws_stats.cell(row=row_num, column=2, value=bulgu_sayisi)
        ws_stats.cell(row=row_num, column=3, value=uniq_zafiyet_sayisi)
        row_num += 1

    # Artık istatistik dosyasını kaydedelim
    try:
        wb_stats.save(istatistik_output_excel_path)
        print(Fore.CYAN + Style.BRIGHT + f"[INFO] İstatistik dosyası oluşturuldu: {istatistik_output_excel_path}" + Style.RESET_ALL)
    except Exception as e:
        print(Fore.RED + f"[-] Hata: İstatistik Excel kaydedilemedi.\n{e}" + Style.RESET_ALL)

    # Özet
    print(Fore.GREEN + f"\n[INFO] Toplam {total_vulns} adet zafiyet işlendi (filtreye uygunsa)." + Style.RESET_ALL)


def main():
    """
    Kullanıcıyla etkileşimli olarak .nessus dosyası adı, çıktı dosyası adı 
    ve risk seviyesi filtrelerini alır.
    Ardından parse_nessus_to_excel fonksiyonunu çağırır.
    """
    print(Fore.GREEN + Style.BRIGHT + "***** Nessus Parser Programı *****" + Style.RESET_ALL)
    print(Fore.GREEN + "Bu program, .nessus raporunu parse ederek iki Excel çıktısı almanıza yardımcı olur.\n" + Style.RESET_ALL)

    # 1) Nessus dosyasının adını sor
    nessus_file_path = input(Fore.BLUE + "Lütfen .nessus dosyasının yolunu giriniz (örn: rapor.nessus): " + Style.RESET_ALL).strip()
    if not nessus_file_path:
        print(Fore.RED + "[-] Hata: Nessus dosyası adı boş olamaz!" + Style.RESET_ALL)
        return

    # 2) Çıktı dosyası adı
    output_excel_path = input(Fore.BLUE + "Oluşturulacak Excel dosyasının adı ne olsun? (örn: cikti.xlsx) [varsayılan: nessus_output.xlsx]: " + Style.RESET_ALL).strip()
    if not output_excel_path:
        output_excel_path = "nessus_output.xlsx"

    # 3) Hangi risk seviyeleri (filtre)?
    all_severities = ["None", "Info", "Low", "Medium", "High", "Critical"]
    print(Fore.BLUE + "Desteklenen zafiyet seviyeleri: " + Style.RESET_ALL + ", ".join(all_severities))
    print(Fore.MAGENTA + "Örneğin: High,Medium ya da sadece Critical gibi. Boş bırakırsanız hepsi dahil olur." + Style.RESET_ALL)
    
    severities_str = input(Fore.BLUE + "Hangi seviyeler rapora dahil edilsin? (virgülle ayrılmış): " + Style.RESET_ALL).strip()
    if severities_str:
        severities_list = [sev.strip() for sev in severities_str.split(",") if sev.strip()]
    else:
        severities_list = None

    # Çalıştır
    parse_nessus_to_excel(
        nessus_file_path=nessus_file_path,
        output_excel_path=output_excel_path,
        severities=severities_list
    )


if __name__ == "__main__":
    main()
